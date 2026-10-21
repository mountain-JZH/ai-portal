"""Fixed A:AG workbook contract and conservative, lossless project semantics."""

from datetime import date, datetime
from io import BytesIO
import json
import re
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook


STAGES = ["客户对接", "实地勘察", "资料收集与方案设计", "商务洽谈", "可研编制",
          "项目审批", "合同签约", "项目备案", "招标", "施工", "竣工投运"]
FIELDS = ["serial", "leader", "owner", "resource", "name", "management", "ownership",
          "priority", "valueRating", "customerType", "district", "pvArea", "pvCapacity",
          "storageCapacity", "chargingScale", "energyType", "energyScale"]
TAIL_FIELDS = ["currentStage", "situation", "nextStep", "coordination", "progressUpdate"]
BUSINESS_FIELDS = [
    ("pvArea", "光伏可利用面积", "㎡"), ("pvCapacity", "光伏装机规模", "MW"),
    ("storageCapacity", "储能装机容量（原表数值 / 单位）", ""), ("chargingScale", "充电桩规模", ""),
    ("energyType", "分布式能源类型", ""), ("energyScale", "分布式能源规模", ""),
]
PRIORITY_LABELS = ("尽快推进",)  # Explicit business labels; never infer from A/甲 alone.
MAX_FILE_SIZE = 4 * 1024 * 1024
MAX_ROWS = 10000


def init_project_tables(connection):
    # Additive only, inside the existing database initialization transaction.
    statements = [
        """CREATE TABLE IF NOT EXISTS projects (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name_key TEXT NOT NULL UNIQUE,
            resource_key TEXT NOT NULL,
            data TEXT NOT NULL,
            imported_at TEXT NOT NULL,
            import_id INTEGER NOT NULL
        )""",
        """CREATE TABLE IF NOT EXISTS project_imports (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            imported_at TEXT NOT NULL,
            admin_id INTEGER NOT NULL,
            counts TEXT NOT NULL
        )""",
        """CREATE TABLE IF NOT EXISTS project_import_previews (
            token TEXT PRIMARY KEY,
            admin_id INTEGER NOT NULL,
            expires_at INTEGER NOT NULL,
            base_revision TEXT NOT NULL,
            payload TEXT NOT NULL
        )""",
    ]
    for statement in statements:
        connection.execute(statement)
    from project_changes import migrate
    migrate(connection)


def text(value):
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        value = value.isoformat()
    elif isinstance(value, float) and value.is_integer():
        value = int(value)
    result = str(value)
    return result if result.strip() not in {"", "/", "／"} else None


def key(value):
    return re.sub(r"\s+", "", value or "").casefold()


def needs_coordination(value):
    return bool(text(value)) and (value or "").strip() not in {"无", "暂无", "无。", "暂无。"}


def normalize_stage(value):
    raw = (value or "").strip()
    return STAGES[2] if raw in {"资料收集&方案设计", "资料收集＆方案设计"} else raw


def stage_group(project):
    raw = normalize_stage(project.get("currentStage"))
    return raw if raw in STAGES else "未明确"


def enrich(project):
    result = dict(project)
    result["stageGroup"] = stage_group(project)
    result["isPriority"] = any(label in (project.get("priority") or "") for label in PRIORITY_LABELS)
    result["needsCoordination"] = needs_coordination(project.get("coordination"))
    result["isOperational"] = project["stages"][-1] == "已完成"
    result["completedCount"] = sum(status == "已完成" for status in project["stages"])
    return result


def canonical(project):
    # Sequence numbers may change when Excel is sorted; they are display-only.
    return json.dumps({k: v for k, v in project.items() if k != "serial"},
                      ensure_ascii=False, sort_keys=True)


def parse_workbook(content):
    """No formulas/macros/links are executed. Formula cells require Excel cached values."""
    if len(content) > MAX_FILE_SIZE:
        raise ValueError("文件超过 4 MB，请缩小台账后重试")
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 2500 or sum(item.file_size for item in entries) > 32 * 1024 * 1024:
                raise ValueError("Excel 解压后过大，最多支持 32 MB / 2500 个内部文件")
            names = [item.filename.lower() for item in entries]
            if len(set(names)) != len(names) or any(item.flag_bits & 1 for item in entries):
                raise ValueError("不支持加密或重复内部文件的 Excel")
            if any("vbaproject" in name or "externallinks/" in name or
                   name.endswith("connections.xml") for name in names):
                raise ValueError("不支持包含宏、外部工作簿链接或数据连接的文件")
            if "xl/workbook.xml" not in names or "[content_types].xml" not in names:
                raise ValueError("文件不是有效的 .xlsx 工作簿")
    except BadZipFile as error:
        raise ValueError("文件不是有效的 .xlsx 工作簿") from error

    books = []
    try:
        for data_only in (True, False):
            books.append(load_workbook(BytesIO(content), read_only=True,
                                       data_only=data_only, keep_links=False, keep_vba=False))
        if any("Sheet1" not in book.sheetnames for book in books):
            raise ValueError("找不到 Sheet1，请使用原始台账结构")
        sheet, formulas = (book["Sheet1"] for book in books)
        if sheet.max_row > MAX_ROWS + 3 or sheet.max_column > 100:
            raise ValueError("工作表范围过大，最多支持 10000 条数据行")
        headers = list(sheet.iter_rows(min_row=1, max_row=3, max_col=33, values_only=True))
        expected = {2: "负责人", 3: "资源名称", 4: "项目名称", 10: "行政片区",
                    28: "阶段", **{17 + i: stage for i, stage in enumerate(STAGES)}}
        for column, label in expected.items():
            combined = key("".join(str(row[column] or "") for row in headers))
            # A controlled alias for this one heading; column/order validation remains intact.
            if column == 19:
                combined = combined.replace("资料收集&方案设计", key(STAGES[2])).replace("资料收集＆方案设计", key(STAGES[2]))
            if key(label) not in combined:
                raise ValueError(f"无法识别表头：第 {column + 1} 列应包含“{label}”，请保留前三行与 A～AG 列顺序")
        projects, warnings, errors = [], [], []
        skipped = 0
        cached_rows = sheet.iter_rows(min_row=4, max_col=33)
        formula_rows = formulas.iter_rows(min_row=4, max_col=33)
        for row_number, (cells, raw_cells) in enumerate(zip(cached_rows, formula_rows), 4):
            values = [text(cell.value) for cell in cells]
            if not any(values[1:]) and not any(cell.data_type == "f" for cell in raw_cells[1:]):
                skipped += 1
                continue
            for index, cell in enumerate(raw_cells):
                if cell.data_type == "f":
                    if cells[index].value is None:
                        errors.append(f"第 {row_number} 行 {cell.column_letter} 列公式没有缓存值，请用 Excel 重算并保存")
                    else:
                        warnings.append(f"第 {row_number} 行 {cell.column_letter} 列使用 Excel 最近保存的公式结果；请确认已重算")
                if cells[index].data_type == "e":
                    errors.append(f"第 {row_number} 行 {cell.column_letter} 列为 Excel 错误值")
            if not values[4]:
                errors.append(f"第 {row_number} 行缺少项目名称；有业务内容的行不能作为空白行跳过")
                continue
            project = dict(zip(FIELDS, values[:17]))
            project.update(zip(TAIL_FIELDS, values[28:33]))
            project["stages"] = values[17:28]
            for index, status in enumerate(project["stages"]):
                if status is not None and status not in {"已完成", "进行中"}:
                    warnings.append(f"第 {row_number} 行「{STAGES[index]}」状态“{status}”未识别，原值将保留")
            current = normalize_stage(project["currentStage"])
            active = [i for i, status in enumerate(project["stages"]) if status == "进行中"]
            if current not in STAGES:
                warnings.append(f"第 {row_number} 行当前阶段缺失或无法匹配，将计入“未明确”")
            elif ((active and active != [STAGES.index(current)]) or
                  any(s == "已完成" for s in project["stages"][STAGES.index(current) + 1:]) or
                  (current != STAGES[-1] and project["stages"][STAGES.index(current)] == "已完成")):
                warnings.append(f"第 {row_number} 行当前阶段与节点状态不一致，仍保留 AC 原值“{current}”")
            projects.append({"row": row_number, "project": project})
        if not projects and not errors:
            errors.append("没有有效项目，请填写 E 列项目名称；仅有序号的行不计入项目")
        return projects, skipped, warnings, errors
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("Excel 无法安全解析，请检查文件是否损坏并另存为 .xlsx") from error
    finally:
        for book in books:
            book.close()


def plan_import(incoming, existing):
    names = {}
    for row in existing:
        names.setdefault(key(row["name"]), []).append(row)
    seen_names, seen_resources = set(), set()
    counts = {"valid": len(incoming), "new": 0, "updated": 0, "unchanged": 0}
    operations, errors = [], []
    for item in incoming:
        project = item["project"]
        name, resource = key(project["name"]), key(project["resource"])
        if name in seen_names or (resource and resource in seen_resources):
            errors.append(f"第 {item['row']} 行项目名称或资源名称重复，无法唯一匹配，请在台账中消除冲突")
        seen_names.add(name)
        if resource:
            seen_resources.add(resource)
        matches = names.get(name, [])
        resource_matches = [row for row in existing if resource and key(row["resource"]) == resource]
        if len(matches) > 1 or (matches and key(matches[0]["resource"]) != resource) or (
                resource_matches and (len(resource_matches) != 1 or key(resource_matches[0]["name"]) != name)):
            errors.append(f"第 {item['row']} 行项目关键字段变化或匹配冲突，已阻止整批写入；请核对项目名称与资源名称")
            continue
        match = matches[0] if matches else None
        action = "new" if match is None else "unchanged" if canonical(project) == canonical(
            {field: match.get(field) for field in project}) else "updated"
        counts[action] += 1
        operations.append({**item, "id": match["id"] if match else None, "action": action})
    return counts, operations, errors
