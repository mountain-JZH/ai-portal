"""Project import contract and authenticated HTTP tests; only temporary data is used."""

import asyncio
from contextlib import closing
from http.cookies import SimpleCookie
from io import BytesIO
import json
import os
from pathlib import Path
import sqlite3
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree

from openpyxl import Workbook
import requests

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from test_section_visuals import call_app
from project_data import FIELDS, STAGES


def sample_project(**changes):
    project = dict.fromkeys(FIELDS)
    project.update({"serial": "1", "leader": "组长A", "owner": "负责人B", "resource": "XX希尔顿酒店",
                    "name": "XX希尔顿酒店分布式光伏项目", "priority": "甲：尽快推进", "valueRating": "A级",
                    "district": "南京市江宁区", "pvArea": "15000㎡", "pvCapacity": "1.8MW",
                    "stages": ["已完成"] * 3 + ["进行中"] + [None] * 7,
                    "currentStage": "商务洽谈",
                    "situation": "管委会不同意建设光伏，理由还是影响酒店美观度，保持跟进。",
                    "nextStep": "关注央巡结束时间，完善合作方案，争取合作机会。",
                    "coordination": None, "progressUpdate": "无"})
    project.update(changes)
    return project


def workbook_bytes(projects=None, placeholders=3):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.merge_cells("A1:AG1")
    sheet["A1"] = "项目推进台账 · 按用户字段契约生成的测试夹具，非原始文件"
    sheet.merge_cells("A2:K2")
    sheet["A2"] = "项目基础信息"
    headers = ["序号", "项目组长", "项目负责人", "资源名称", "项目名称", "管理权", "资产所有权",
               "项目开发时序评级（等级）", "项目开发价值评级（星级）", "客户性质", "所属行政片区",
               "分布式光伏预计可利用面积（㎡）", "分布式光伏预计装机规模（MW）", "储能预计装机容量（MW/MWh）",
               "充电桩预计规模", "分布式能源类型", "分布式能源规模", *STAGES,
               "目前所处阶段", "当前具体推进情况、存在的问题、储备原因等", "下一步打算", "需上级协调事项", "是否有进度更新"]
    for column, header in enumerate(headers, 1):
        sheet.cell(3, column, header)
    for project in projects if projects is not None else [sample_project()]:
        sheet.append([project.get(field) for field in FIELDS] + project["stages"] +
                     [project.get(field) for field in ["currentStage", "situation", "nextStep", "coordination", "progressUpdate"]])
    for number in range(placeholders):
        sheet.append([number + 2])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class ProjectTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(prefix="ai-portal-project-tests-")
        cls.root = Path(cls.temp.name)
        cls.env = patch.dict(os.environ, {
            "DATABASE_PATH": str(cls.root / "test.db"), "SESSION_SECRET": "project-tests-only-local",
            "SESSION_COOKIE_SECURE": "false", "ADMIN_USERNAME": "project-test-admin",
            "ADMIN_PASSWORD": "project-test-password",
        })
        cls.env.start()
        import database
        import upload_api
        cls.database = database
        cls.db_patch = patch.object(database, "DATABASE_PATH", cls.root / "test.db")
        cls.upload_patch = patch.object(upload_api, "UPLOAD_ROOT", cls.root / "uploads")
        cls.db_patch.start()
        cls.upload_patch.start()
        import main
        cls.app = main.app
        from auth_support import employee_session
        cls.employee_cookie, cls.csrf = employee_session(cls.app, call_app)
        status, headers, login_result = cls.request("POST", "/api/auth/login", {
            "username": "project-test-admin", "password": "project-test-password"}, admin=False)
        assert status == 200
        cookies = SimpleCookie()
        cookies.load(headers.get(b"set-cookie", cls.employee_cookie.encode()).decode())
        cls.cookie = "; ".join(f"{key}={value.value}" for key, value in cookies.items())
        cls.admin_csrf = login_result['csrfToken']
        cls.employee_cookie, cls.csrf = employee_session(cls.app, call_app)

    @classmethod
    def tearDownClass(cls):
        cls.upload_patch.stop()
        cls.db_patch.stop()
        cls.env.stop()
        cls.temp.cleanup()

    @classmethod
    def request(cls, method, path, payload=None, admin=True, query=""):
        headers, body = {"cookie": cls.employee_cookie, "x-csrf-token": cls.csrf}, b""
        if payload is not None:
            body = json.dumps(payload).encode()
            headers["content-type"] = "application/json"
        if admin:
            headers["cookie"] = cls.cookie
            headers["x-csrf-token"] = cls.admin_csrf
        app = cls.app
        if query:
            async def queried(scope, receive, send):
                scope["query_string"] = query.encode()
                await cls.app(scope, receive, send)
            app = queried
        status, returned_headers, content = asyncio.run(call_app(app, method, path, body, headers))
        return status, returned_headers, json.loads(content)

    def upload(self, content=None, admin=True, filename="台账.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"):
        prepared = requests.Request("POST", "http://testserver", files={
            "file": (filename, workbook_bytes() if content is None else content, mime),
        }).prepare()
        headers = {"content-type": prepared.headers["Content-Type"], "cookie": self.employee_cookie, "x-csrf-token": self.csrf}
        if admin:
            headers["cookie"] = self.cookie
            headers["x-csrf-token"] = self.admin_csrf
        status, _, body = asyncio.run(call_app(self.app, "POST", "/api/admin/projects/imports/preview", prepared.body, headers))
        return status, json.loads(body)

    def confirm(self, token, **extra):
        status, _, data = self.request("POST", "/api/admin/projects/imports/confirm", {"token": token, **extra})
        return status, data

    def import_projects(self, projects=None):
        status, preview = self.upload(workbook_bytes(projects))
        self.assertEqual(status, 200, preview)
        self.assertTrue(preview["canConfirm"], preview)
        status, result = self.confirm(preview["token"])
        self.assertEqual(status, 200, result)
        return result

    def dashboard(self, **filters):
        from urllib.parse import urlencode
        status, _, data = self.request("GET", "/api/admin/projects", query=urlencode(filters))
        self.assertEqual(status, 200)
        return data

    def setUp(self):
        with closing(self.database.get_db_connection()) as connection, connection:
            for table in ["projects", "project_imports", "project_import_previews"]:
                connection.execute(f"DELETE FROM {table}")

    def test_first_import_sample_detail_and_real_summary(self):
        status, preview = self.upload()
        self.assertEqual(status, 200)
        self.assertEqual(preview["counts"], {"valid": 1, "new": 1, "updated": 0, "unchanged": 0, "invalid": 3, "skipped": 3})
        self.assertEqual(self.dashboard()["summary"]["total"], 0, "Preview must not write projects")
        self.assertEqual(self.confirm(preview["token"])[0], 200)
        result = self.dashboard()
        self.assertEqual(result["summary"], {"total": 1, "priority": 1, "coordination": 0, "operational": 0})
        self.assertEqual(sum(s["count"] for s in result["distribution"]), 1)
        self.assertEqual(next(s["count"] for s in result["distribution"] if s["name"] == "商务洽谈"), 1)
        project = result["projects"][0]
        self.assertEqual(project["completedCount"], 3)
        self.assertEqual(project["progressUpdate"], "无")
        self.assertTrue(all(value is None for value in project["stages"][4:]))
        status, _, detail = self.request("GET", f"/api/admin/projects/{project['id']}")
        self.assertEqual(status, 200)
        for field, value in sample_project().items():
            self.assertEqual(detail[field], value)
        self.assertTrue(detail["importedAt"])
        self.assertEqual(self.confirm(preview["token"])[0], 409, "Token must be single-use")

    def test_repeat_serial_reorder_update_full_text_and_omitted_old_preserved(self):
        self.import_projects([sample_project(), sample_project(name="独立项目", resource="独立资源", serial="2")])
        status, preview = self.upload(workbook_bytes([sample_project(serial="99")]))
        self.assertEqual(status, 200)
        self.assertEqual(preview["counts"]["unchanged"], 1)
        self.assertEqual(self.confirm(preview["token"])[0], 200)
        full_text = "完整情况\n  继续跟进  " * 400
        status, preview = self.upload(workbook_bytes([sample_project(situation=full_text, coordination="请协调现场踏勘\n保留全文")]))
        self.assertEqual(preview["counts"]["updated"], 1)
        self.assertEqual(self.confirm(preview["token"])[0], 200)
        result = self.dashboard()
        self.assertEqual(result["summary"]["total"], 2)
        self.assertEqual(result["summary"]["coordination"], 1)
        self.assertEqual(result["projects"][0]["situation"], full_text)

    def test_duplicates_name_resource_and_key_changes_block_entire_import(self):
        self.import_projects()
        original = self.dashboard()["projects"]
        for projects in [
            [sample_project(), sample_project()],
            [sample_project(resource="改名资源")],
            [sample_project(name="改名项目")],
            [sample_project(name="另一个", resource="不同"), sample_project(name="新名称", resource="不同")],
        ]:
            status, preview = self.upload(workbook_bytes(projects))
            self.assertEqual(status, 200)
            self.assertFalse(preview["canConfirm"])
            self.assertIsNone(preview["token"])
            self.assertTrue(preview["errors"])
            self.assertEqual(self.dashboard()["projects"], original)

    def test_unknown_status_stage_mismatch_unmatched_stage_retained(self):
        project = sample_project(currentStage="施工", stages=["已完成", "进行中", "待领导确认"] + [None] * 8)
        status, preview = self.upload(workbook_bytes([project]))
        self.assertEqual(status, 200)
        self.assertTrue(any("未识别" in message for message in preview["warnings"]))
        self.assertTrue(any("不一致" in message for message in preview["warnings"]))
        self.assertEqual(self.confirm(preview["token"])[0], 200)
        self.assertEqual(self.dashboard()["projects"][0]["stages"][2], "待领导确认")
        self.import_projects([sample_project(currentStage="意向储备")])
        self.assertEqual(self.dashboard()["projects"][0]["stageGroup"], "未明确")

    def test_coordination_priority_and_operational_are_conservative(self):
        projects = [sample_project(name=f"项目{index}", resource=f"资源{index}", coordination=value,
                                   priority="A级", valueRating="A级", currentStage="竣工投运")
                    for index, value in enumerate([None, "无", "/", "需协调设备接入"])]
        projects[-1]["stages"] = [None] * 10 + ["已完成"]
        self.import_projects(projects)
        self.assertEqual(self.dashboard()["summary"], {"total": 4, "priority": 0, "coordination": 1, "operational": 1})
        self.assertEqual(self.dashboard(focus="coordination")["filteredCount"], 1)

    def test_server_side_search_and_filters(self):
        self.import_projects([sample_project(), sample_project(name="独立项目", resource="独立资源", district="镇江市", currentStage="施工")])
        self.assertEqual(self.dashboard(search="希尔顿")["filteredCount"], 1)
        self.assertEqual(self.dashboard(stage="施工")["projects"][0]["name"], "独立项目")
        self.assertEqual(self.dashboard(district="南京市江宁区")["filteredCount"], 1)
        self.assertEqual(self.dashboard(search="不存在")["filteredCount"], 0)

    def test_missing_name_malformed_headers_and_only_serial_rows(self):
        self.import_projects()
        for content in [workbook_bytes([sample_project(name=None)]), workbook_bytes([])]:
            status, preview = self.upload(content)
            self.assertEqual(status, 200)
            self.assertFalse(preview["canConfirm"])
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(workbook_bytes()))
        wb.active["E3"] = "错误表头"
        output = BytesIO()
        wb.save(output)
        wb.close()
        self.assertEqual(self.upload(output.getvalue())[0], 400)
        self.assertEqual(self.dashboard()["summary"]["total"], 1)

    def test_formulas_numeric_units_and_placeholders(self):
        content = workbook_bytes([sample_project(pvArea=15000, pvCapacity=1.8, storageCapacity="/", nextStep="=1+1")])
        status, preview = self.upload(content)
        self.assertEqual(status, 200)
        self.assertFalse(preview["canConfirm"])
        self.assertTrue(any("没有缓存值" in message for message in preview["errors"]))
        # Inject Excel's cached result into the test fixture, without executing a formula.
        original = ZipFile(BytesIO(content))
        output = BytesIO()
        with original, ZipFile(output, "w", ZIP_DEFLATED) as archive:
            for entry in original.infolist():
                value = original.read(entry.filename)
                if entry.filename == "xl/worksheets/sheet1.xml":
                    tree = ElementTree.fromstring(value)
                    namespace = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                    cached = tree.find(".//m:c[@r='AE4']/m:v", namespace)
                    self.assertIsNotNone(cached)
                    cached.text = "2"
                    value = ElementTree.tostring(tree)
                archive.writestr(entry, value)
        status, preview = self.upload(output.getvalue())
        self.assertEqual(status, 200)
        self.assertTrue(preview["canConfirm"], preview)
        self.assertEqual(self.confirm(preview["token"])[0], 200)
        project = self.dashboard()["projects"][0]
        self.assertEqual(project["pvArea"], "15000")
        self.assertEqual(project["pvCapacity"], "1.8")
        self.assertIsNone(project["storageCapacity"])
        self.assertEqual(project["nextStep"], "2")

    def test_auth_all_reads_writes_nonadmin_and_inactive(self):
        for method, path, payload in [("GET", "/api/admin/projects", None),
                                     ("GET", "/api/admin/projects/1", None),
                                     ("POST", "/api/admin/projects/imports/confirm", {"token": "x" * 40}),
                                     ("DELETE", "/api/admin/projects/imports/preview/unknown", None)]:
            self.assertEqual(self.request(method, path, payload, admin=False)[0], 403)
        self.assertEqual(self.upload(admin=False)[0], 403)
        self.assertEqual(self.request("POST", "/api/auth/login", {"username": "reader", "password": "not-an-admin"}, admin=False)[0], 401)
        with closing(self.database.get_db_connection()) as connection, connection:
            connection.execute("UPDATE admins SET is_active = 0")
        try:
            self.assertEqual(self.upload()[0], 401)
            self.assertEqual(self.request("GET", "/api/admin/projects")[0], 401)
        finally:
            with closing(self.database.get_db_connection()) as connection, connection:
                connection.execute("UPDATE admins SET is_active = 1")

    def test_tampered_expired_cancelled_wrong_owner_and_stale_tokens(self):
        _, preview = self.upload()
        self.assertEqual(self.confirm(preview["token"], operations=[])[0], 422)
        self.assertEqual(self.confirm("invalid-token" * 4)[0], 409)
        with closing(self.database.get_db_connection()) as connection, connection:
            connection.execute("UPDATE project_import_previews SET expires_at = 0")
        self.assertEqual(self.confirm(preview["token"])[0], 409)
        _, preview = self.upload()
        self.assertEqual(self.request("DELETE", f"/api/admin/projects/imports/preview/{preview['token']}")[0], 200)
        self.assertEqual(self.confirm(preview["token"])[0], 409)
        _, preview = self.upload()
        with closing(self.database.get_db_connection()) as connection, connection:
            connection.execute("UPDATE project_import_previews SET admin_id = -1")
        self.assertEqual(self.confirm(preview["token"])[0], 409)
        _, preview = self.upload()
        with closing(self.database.get_db_connection()) as connection, connection:
            connection.execute("INSERT INTO project_imports (imported_at, admin_id, counts) VALUES ('test', 1, '{}')")
        self.assertEqual(self.confirm(preview["token"])[0], 409)
        self.assertEqual(self.dashboard()["summary"]["total"], 0)

    def test_mid_batch_database_failure_rolls_back_everything(self):
        self.import_projects()
        original = self.dashboard()
        _, preview = self.upload(workbook_bytes([sample_project(situation="试图更新"), sample_project(name="触发失败", resource="新资源")]))
        with closing(self.database.get_db_connection()) as connection, connection:
            connection.execute("CREATE TRIGGER project_test_failure BEFORE INSERT ON projects BEGIN SELECT RAISE(ABORT, 'injected test failure'); END")
        try:
            with self.assertRaises(sqlite3.IntegrityError):
                self.confirm(preview["token"])
        finally:
            with closing(self.database.get_db_connection()) as connection, connection:
                connection.execute("DROP TRIGGER project_test_failure")
        after = self.dashboard()
        self.assertEqual(original["projects"], after["projects"])
        self.assertEqual(original["latestImport"], after["latestImport"])

    def test_file_type_corruption_macros_connections_zip_bomb_and_size(self):
        self.assertEqual(self.upload(filename="台账.xlsm")[0], 400)
        self.assertEqual(self.upload(mime="text/plain")[0], 400)
        self.assertEqual(self.upload(b"not-an-xlsx")[0], 400)
        self.assertEqual(self.upload(b"x" * (4 * 1024 * 1024 + 1))[0], 400)
        self.assertEqual(self.upload(b"x" * (4 * 1024 * 1024 + 70 * 1024))[0], 413)
        for filename, content in [("xl/vbaProject.bin", b"test"), ("xl/connections.xml", b"test"),
                                  ("xl/externalLinks/externalLink1.xml", b"test"),
                                  ("large.xml", b"x" * (33 * 1024 * 1024))]:
            output = BytesIO()
            with ZipFile(BytesIO(workbook_bytes())) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
                for entry in source.infolist():
                    target.writestr(entry, source.read(entry.filename))
                target.writestr(filename, content)
            self.assertEqual(self.upload(output.getvalue())[0], 400)

    def test_employee_read_access_empty_detail_and_admin_boundary(self):
        self.assertEqual(self.request("GET", "/api/projects", admin=False)[0], 200)
        self.assertEqual(self.request("GET", "/api/admin/projects", admin=False)[0], 403)
        self.assertEqual(self.upload(admin=False)[0], 403)
        status, _, body = asyncio.run(call_app(self.app, "GET", "/api/projects", b"", {}))
        self.assertEqual(status, 401)
        self.assertNotIn(b'projects', body)
        self.assertEqual(self.request("GET", "/api/projects/9999999", admin=False)[0], 404)
        self.assertEqual(self.request("GET", "/api/projects/no-number", admin=False)[0], 422)
        self.assertEqual(self.request("GET", "/api/projects", admin=False, query="view=unknown")[0], 422)
        self.assertEqual(self.request("GET", "/api/projects", admin=False, query="business=unknown")[0], 422)
        with closing(self.database.get_db_connection()) as connection, connection:
            connection.execute("DELETE FROM projects")
        self.assertEqual(self.request("GET", "/api/projects", admin=False)[2]["projects"], [])

    def test_employee_perspectives_isolation_or_dedup_raw_values_and_alias(self):
        from urllib.parse import urlencode
        from openpyxl import load_workbook
        from project_data import parse_workbook
        fixture = [sample_project(name="虚构测试工程一", resource="虚构资源一", progressUpdate="是",
                                  owner="虚构负责人甲", pvCapacity="0.06699999999999999",
                                  storageCapacity="尚不能确定", energyScale="完整方案说明" * 150),
                   sample_project(name="虚构测试工程二", resource="虚构资源二", progressUpdate="无",
                                  pvArea=None, pvCapacity=None, chargingScale="两台设备，参数待核实"),
                   sample_project(name="虚构测试工程三", resource="虚构资源三", progressUpdate="已更新",
                                  pvArea=None, pvCapacity=None, storageCapacity="7.92E-2")]
        with closing(self.database.get_db_connection()) as connection, connection:
            connection.execute("DELETE FROM projects")
        workbook = load_workbook(BytesIO(workbook_bytes(fixture)))
        workbook.active.cell(3, 20, "资料收集&方案设计")
        output = BytesIO(); workbook.save(output); workbook.close()
        parsed, _, _, errors = parse_workbook(output.getvalue())
        self.assertFalse(errors)
        self.assertEqual(parsed[0]["project"]["pvCapacity"], "0.06699999999999999")
        status, preview = self.upload(output.getvalue())
        self.assertEqual(status, 200)
        self.assertEqual(self.confirm(preview["token"])[0], 200)
        def read(**filters):
            status, _, result = self.request("GET", "/api/projects", admin=False, query=urlencode(filters, doseq=True))
            self.assertEqual(status, 200)
            return result
        self.assertEqual(read(view="overview", search="不存在的测试搜索")["filteredCount"], 0)
        self.assertEqual(read(view="daily", search="不存在的测试搜索", district="不存在的片区")["filteredCount"], 0)
        self.assertEqual(read(view="people", role="owner", person="虚构负责人甲", search="不存在的测试搜索")["filteredCount"], 1)
        businesses = read(view="business", business=["光伏", "储能"], owner="不存在")
        self.assertEqual(businesses["filteredCount"], 2)
        self.assertEqual(len({p["id"] for p in businesses["projects"]}), 2)
        first = next(p for p in businesses["projects"] if p["name"] == fixture[0]["name"])
        detail = self.request("GET", f'/api/projects/{first["id"]}', admin=False)[2]
        self.assertEqual(first, detail)
        self.assertEqual(detail["energyScale"], fixture[0]["energyScale"])
        self.assertEqual(detail["storageCapacity"], "尚不能确定")
        self.assertEqual(detail["completedCount"], 3)
        self.assertEqual(len(detail["stages"]), 11)
        workbook = load_workbook(BytesIO(output.getvalue()))
        workbook.active.cell(3, 20, "错误结构")
        wrong = BytesIO(); workbook.save(wrong); workbook.close()
        self.assertEqual(self.upload(wrong.getvalue())[0], 400)

    def test_eight_invented_acceptance_scenarios_persist_and_read_all_details(self):
        from project_board_fixtures import synthetic_projects
        from urllib.parse import urlencode
        fixtures = synthetic_projects()
        self.import_projects(fixtures)
        status, _, result = self.request("GET", "/api/projects", admin=False)
        self.assertEqual(status, 200)
        self.assertEqual(result["summary"]["total"], 8)
        self.assertEqual(sum(result["personCounts"]["owner"].values()), 8)
        self.assertEqual(result["businessCounts"]["光伏"], 4)
        for project in result["projects"]:
            detail = self.request("GET", f'/api/projects/{project["id"]}', admin=False)[2]
            self.assertEqual(detail, project)
            original = next(item for item in fixtures if item["name"] == project["name"])
            for field in [*FIELDS, "stages", "situation", "nextStep", "coordination", "progressUpdate"]:
                self.assertEqual(detail[field], original[field])
        for view in ["daily", "people", "business"]:
            query = {"view": view, "search": "不存在的隐藏筛选", "district": "不存在的片区"}
            if view == "business": query["business"] = ["光伏", "储能"]
            response = self.request("GET", "/api/projects", admin=False, query=urlencode(query, doseq=True))[2]
            self.assertEqual(response["filteredCount"], {"daily": 0, "people": 8, "business": 5}[view])
            self.assertEqual(len({p["id"] for p in response["projects"]}), response["filteredCount"])

    def test_idempotent_migrations_preserve_existing_content_and_reconnect(self):
        with closing(self.database.get_db_connection()) as connection, connection:
            connection.execute("INSERT INTO news (title) VALUES ('项目迁移不改变新闻')")
        self.import_projects()
        before = self.dashboard()["projects"]
        self.database.init_database()
        self.database.init_database()
        self.assertEqual(self.dashboard()["projects"], before)
        with closing(self.database.get_db_connection()) as connection:
            self.assertEqual(connection.execute("SELECT title FROM news WHERE title = '项目迁移不改变新闻'").fetchone()[0], "项目迁移不改变新闻")


if __name__ == "__main__":
    unittest.main(verbosity=2)
